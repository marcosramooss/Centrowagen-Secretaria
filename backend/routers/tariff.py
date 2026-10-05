"""Tariff import — sube el Excel/CSV oficial, revisa el resumen y aplica.

Two steps on purpose: /tariff/preview never writes, /tariff/apply writes only
the rows the user confirmed. Matching is by modelo + acabado.
"""

import csv
import io
import logging
import re
import unicodedata

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile

from lib.auth import require_admin
from lib.db import db
from lib.doc import new_id, now_utc
from models.tariff import TariffApplyIn, TariffApplyResult, TariffPreview, TariffRow

router = APIRouter(tags=["tariff"])
logger = logging.getLogger(__name__)

# logical field → accepted header spellings (accent/case-insensitive)
HEADERS = {
    "model": ["modelo", "model", "vehiculo", "vehículo", "gama"],
    "trim": ["acabado", "version", "versión", "trim", "equipamiento"],
    "base_price": ["pvp", "precio", "precio base", "pvp recomendado", "base_price", "tarifa"],
    "promotional_price": ["promocional", "precio promocional", "pvp promocional", "oferta", "precio oferta"],
    "financing_price": ["financiado", "precio financiado", "con financiacion", "con financiación", "pvp financiado"],
    "discount": ["descuento", "dto", "dto.", "descuento %", "%"],
    "campaign": ["campana", "campaña", "campaign", "promocion", "promoción"],
}

MAX_BYTES = 10 * 1024 * 1024


def norm(s: object) -> str:
    text = str(s or "").strip().lower()
    text = unicodedata.normalize("NFKD", text)
    return "".join(c for c in text if not unicodedata.combining(c))


def to_number(value: object) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    raw = str(value).strip()
    if not raw:
        return None
    raw = re.sub(r"[^\d,.\-]", "", raw)
    if not raw:
        return None
    # Spanish format: 35.020,50 → 35020.50 ; also tolerate 35020.5
    if "," in raw and "." in raw:
        raw = raw.replace(".", "").replace(",", ".")
    elif "," in raw:
        raw = raw.replace(",", ".")
    elif raw.count(".") == 1 and len(raw.split(".")[1]) == 3:
        raw = raw.replace(".", "")  # 35.020 → 35020
    try:
        return float(raw)
    except ValueError:
        return None


def _read_table(filename: str, data: bytes) -> list[list[object]]:
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext in {"xlsx", "xlsm", "xls"}:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    if ext in {"csv", "txt"}:
        text = data.decode("utf-8-sig", errors="ignore")
        sample = text[:4000]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
            delim = dialect.delimiter
        except Exception:
            delim = ";" if sample.count(";") > sample.count(",") else ","
        return [list(r) for r in csv.reader(io.StringIO(text), delimiter=delim)]
    raise HTTPException(status_code=400, detail="Formato no soportado. Sube un Excel (.xlsx) o un CSV.")


def _map_columns(rows: list[list[object]]) -> tuple[dict[str, int], dict[str, str], int]:
    """Find the header row anywhere in the first 10 lines (tariffs often have a title block)."""
    for idx, row in enumerate(rows[:10]):
        cells = [norm(c) for c in row]
        mapping: dict[str, int] = {}
        found: dict[str, str] = {}
        for field, names in HEADERS.items():
            for ci, cell in enumerate(cells):
                if not cell or ci in mapping.values():
                    continue
                if cell in names or any(cell.startswith(n) for n in names):
                    mapping[field] = ci
                    found[field] = str(row[ci])
                    break
        if "model" in mapping and "base_price" in mapping:
            return mapping, found, idx
    raise HTTPException(
        status_code=400,
        detail="No encuentro las columnas mínimas. El archivo debe tener una cabecera con 'Modelo' y 'PVP'.",
    )


@router.post("/tariff/preview", response_model=TariffPreview)
async def preview_tariff(file: UploadFile = File(...), _: dict = Depends(require_admin)):
    data = await file.read()
    if len(data) > MAX_BYTES:
        raise HTTPException(status_code=413, detail="El archivo supera el límite de 10 MB")
    if not data:
        raise HTTPException(status_code=400, detail="El archivo está vacío")

    table = _read_table(file.filename or "", data)
    mapping, found, header_idx = _map_columns(table)

    vehicles = await db.vehicles.find({"active": True}).to_list(500)
    prices = await db.prices.find().to_list(2000)
    price_by_vehicle = {p["vehicle_id"]: p for p in prices}

    def cell(row: list[object], field: str) -> object:
        ci = mapping.get(field)
        return row[ci] if ci is not None and ci < len(row) else None

    out: list[TariffRow] = []
    for offset, raw in enumerate(table[header_idx + 1 :], start=header_idx + 2):
        if not raw or all(c is None or str(c).strip() == "" for c in raw):
            continue
        model = str(cell(raw, "model") or "").strip()
        trim = str(cell(raw, "trim") or "").strip()
        base = to_number(cell(raw, "base_price"))

        item = TariffRow(
            row=offset,
            model=model,
            trim=trim,
            base_price=base,
            promotional_price=to_number(cell(raw, "promotional_price")),
            financing_price=to_number(cell(raw, "financing_price")),
            discount=to_number(cell(raw, "discount")),
            campaign=(str(cell(raw, "campaign")).strip() or None) if cell(raw, "campaign") else None,
            status="no_reconocida",
        )

        if not model or base is None or base <= 0:
            item.message = "Falta el modelo o el PVP es inválido"
            out.append(item)
            continue

        match = next(
            (v for v in vehicles if norm(v.get("model")) == norm(model) and norm(v.get("trim")) == norm(trim)),
            None,
        )
        if match is None and not trim:
            candidates = [v for v in vehicles if norm(v.get("model")) == norm(model)]
            if len(candidates) == 1:
                match = candidates[0]

        if match:
            current = price_by_vehicle.get(match["id"], {}).get("base_price")
            item.status = "actualizable"
            item.vehicle_id = match["id"]
            item.current_price = current
            item.trim = item.trim or str(match.get("trim") or "")
            if current is not None and abs(float(current) - base) < 0.01:
                item.message = "Sin cambios (mismo PVP)"
            elif current is not None:
                diff = base - float(current)
                item.message = f"{'+' if diff > 0 else ''}{diff:,.0f} € respecto a la tarifa actual".replace(",", ".")
            else:
                item.message = "Se creará la tarifa de este vehículo"
        else:
            item.status = "nueva"
            item.message = "No existe este modelo/acabado en el catálogo: se ignorará salvo que lo des de alta"
        out.append(item)

    return TariffPreview(
        filename=file.filename or "tarifa",
        detected_columns=found,
        total_rows=len(out),
        updatable=sum(1 for r in out if r.status == "actualizable"),
        new_rows=sum(1 for r in out if r.status == "nueva"),
        unrecognized=sum(1 for r in out if r.status == "no_reconocida"),
        rows=out,
    )


@router.post("/tariff/apply", response_model=TariffApplyResult)
async def apply_tariff(payload: TariffApplyIn, admin: dict = Depends(require_admin)):
    updated = created = stock_updated = skipped = 0
    messages: list[str] = []
    now = now_utc()

    for row in payload.rows:
        # Only confirmed, matched rows are written — never a "nueva"/"no_reconocida".
        if row.status != "actualizable" or not row.vehicle_id or row.base_price is None:
            skipped += 1
            continue
        if not await db.vehicles.find_one({"id": row.vehicle_id}):
            skipped += 1
            messages.append(f"Fila {row.row}: el vehículo ya no existe")
            continue

        fields = {
            "base_price": row.base_price,
            "promotional_price": row.promotional_price,
            "financing_price": row.financing_price,
            "discount": row.discount,
            "campaign": row.campaign,
            "source": payload.source,
            "last_updated": now,
        }
        existing = await db.prices.find_one({"vehicle_id": row.vehicle_id})
        if existing:
            await db.prices.update_one({"vehicle_id": row.vehicle_id}, {"$set": fields})
            updated += 1
        else:
            await db.prices.insert_one({"id": new_id(), "vehicle_id": row.vehicle_id, **fields})
            created += 1

        if payload.update_stock:
            res = await db.stock.update_many(
                {"vehicle_id": row.vehicle_id},
                {"$set": {"pvp": row.base_price, "promotional_price": row.promotional_price, "last_updated": now}},
            )
            stock_updated += res.modified_count

    messages.insert(
        0,
        f"Tarifa aplicada por {admin.get('name')}: {updated} actualizadas, {created} creadas, "
        f"{stock_updated} unidades de stock recalculadas, {skipped} filas ignoradas.",
    )
    return TariffApplyResult(
        prices_updated=updated,
        prices_created=created,
        stock_updated=stock_updated,
        skipped=skipped,
        messages=messages,
    )
