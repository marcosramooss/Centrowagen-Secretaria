"""Knowledge plane: documents (RAG base), FAQ, memory, argumentario."""

import io
import uuid

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, Response, UploadFile

from lib.auth import get_current_user, require_admin
from lib.db import db
from lib.doc import new_id, now_utc, prepare
from lib.rag import score, tokens
from lib.storage import APP_NAME, content_type_for, get_object, put_object
from models.knowledge import (
    Argumentario,
    ArgumentarioIn,
    DocumentItem,
    FaqIn,
    FaqItem,
    MemoryIn,
    MemoryItem,
)

router = APIRouter(tags=["knowledge"])


def _extract_text(filename: str, data: bytes) -> str:
    """Extract plain text for the RAG layer; PDF/DOCX stay metadata-only this phase."""
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    try:
        if suffix in {"txt", "csv", "md", "json"}:
            return data.decode("utf-8", errors="ignore")[:20000]
        if suffix in {"xlsx", "xls"}:
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(data), read_only=True)
            lines: list[str] = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    line = " | ".join("" if c is None else str(c) for c in row)
                    if line.strip(" |"):
                        lines.append(line)
            return "\n".join(lines)[:20000]
        return ""
    except Exception:
        return ""


# ------------------------------------------------------------ Documents


@router.get("/documents", response_model=list[DocumentItem])
async def list_documents(q: str | None = None, category: str | None = None, _: dict = Depends(get_current_user)):
    flt: dict = {"is_deleted": {"$ne": True}}
    if category:
        flt["category"] = category
    docs = await db.documents.find(flt).sort("upload_date", -1).to_list(500)
    items = [DocumentItem(**prepare(d)) for d in docs]
    if q:
        tks = tokens(q)
        scored = [d for d in items if score(tks, d.name, d.category, d.content_text[:4000]) > 0]
        if scored:
            items = scored
    return items


MAX_UPLOAD_BYTES = 25 * 1024 * 1024


@router.post("/documents", response_model=DocumentItem)
async def upload_document(
    name: str = Form(...),
    category: str = Form("General"),
    document_type: str = Form("PDF"),
    source: str = Form("Centrowagen"),
    file: UploadFile | None = File(None),
    admin: dict = Depends(require_admin),
):
    data = await file.read() if file else b""
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="El archivo supera el límite de 25 MB")
    filename = file.filename if file else None
    item = DocumentItem(
        name=name,
        category=category,
        document_type=document_type,
        source=source,
        content_text=_extract_text(filename or "", data),
        size_bytes=len(data) or None,
    )
    storage_path: str | None = None
    if data and filename:
        ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else "bin"
        path = f"{APP_NAME}/documents/{admin['id']}/{uuid.uuid4()}.{ext}"
        try:
            result = await put_object(path, data, content_type_for(filename, file.content_type if file else None))
            storage_path = result["path"]  # always the canonical value returned by storage
            item.file_url = f"/api/documents/{item.id}/file"
        except Exception as exc:
            raise HTTPException(status_code=502, detail=f"No se pudo almacenar el archivo: {exc}")
    doc = item.model_dump()
    doc["storage_path"] = storage_path
    doc["content_type"] = content_type_for(filename or "", file.content_type if file else None) if filename else None
    doc["original_filename"] = filename
    doc["is_deleted"] = False
    await db.documents.insert_one(doc)
    return item


@router.get("/documents/{did}/file")
async def download_document(did: str, authorization: str | None = None, auth: str | None = Query(None), _: dict = Depends(get_current_user)):
    doc = await db.documents.find_one({"id": did, "is_deleted": {"$ne": True}})
    if not doc:
        raise HTTPException(status_code=404, detail="Documento no encontrado")
    if not doc.get("storage_path"):
        raise HTTPException(status_code=404, detail="Este documento no tiene archivo almacenado")
    try:
        data, ctype = await get_object(doc["storage_path"])
    except Exception as exc:
        raise HTTPException(status_code=502, detail=f"No se pudo recuperar el archivo: {exc}")
    return Response(
        content=data,
        media_type=doc.get("content_type") or ctype,
        headers={"Content-Disposition": f'inline; filename="{doc.get("original_filename") or doc.get("name")}"'},
    )


@router.delete("/documents/{did}")
async def delete_document(did: str, _: dict = Depends(require_admin)):
    # Soft delete: object storage has no delete API, the DB is the source of truth.
    res = await db.documents.update_one({"id": did}, {"$set": {"is_deleted": True, "active": False}})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Documento no encontrado")
    return {"ok": True}


# ------------------------------------------------------------ FAQ
# /faq/categories MUST be declared before /faq/{fid} so "categories" is not captured as an id.


@router.get("/faq/categories", response_model=list[str])
async def faq_categories(_: dict = Depends(get_current_user)):
    return sorted({d.get("category", "General") for d in await db.faq.find().to_list(500)})


@router.get("/faq", response_model=list[FaqItem])
async def list_faq(q: str | None = None, category: str | None = None, _: dict = Depends(get_current_user)):
    flt: dict = {}
    if category:
        flt["category"] = category
    docs = await db.faq.find(flt).to_list(500)
    items = [FaqItem(**prepare(d, "last_updated")) for d in docs]
    if q:
        tks = tokens(q)
        scored = [f for f in items if score(tks, f.question, f.answer) > 0]
        if scored:
            items = scored
    return items


@router.post("/faq", response_model=FaqItem)
async def create_faq(payload: FaqIn, _: dict = Depends(require_admin)):
    obj = FaqItem(**payload.model_dump())
    await db.faq.insert_one(obj.model_dump())
    return obj


@router.patch("/faq/{fid}", response_model=FaqItem)
async def update_faq(fid: str, payload: FaqIn, _: dict = Depends(require_admin)):
    if not await db.faq.find_one({"id": fid}):
        raise HTTPException(status_code=404, detail="FAQ no encontrada")
    updates = payload.model_dump(exclude_unset=True)
    updates["last_updated"] = now_utc()
    await db.faq.update_one({"id": fid}, {"$set": updates})
    fresh = await db.faq.find_one({"id": fid})
    return FaqItem(**prepare(fresh, "last_updated"))


@router.delete("/faq/{fid}")
async def delete_faq(fid: str, _: dict = Depends(require_admin)):
    await db.faq.delete_one({"id": fid})
    return {"ok": True}


# ------------------------------------------------------------ Memory


@router.get("/memory", response_model=list[MemoryItem])
async def list_memory(q: str | None = None, category: str | None = None, _: dict = Depends(get_current_user)):
    flt: dict = {}
    if category:
        flt["category"] = category
    docs = await db.memory.find(flt).sort("updated_at", -1).to_list(500)
    items = [MemoryItem(**prepare(d, "created_at", "updated_at")) for d in docs]
    if q:
        tks = tokens(q)
        scored = [m for m in items if score(tks, m.content, m.category) > 0]
        if scored:
            items = scored
    return items


@router.post("/memory", response_model=MemoryItem)
async def create_memory(payload: MemoryIn, _: dict = Depends(require_admin)):
    obj = MemoryItem(**payload.model_dump())
    await db.memory.insert_one(obj.model_dump())
    return obj


@router.patch("/memory/{mid}", response_model=MemoryItem)
async def update_memory(mid: str, payload: MemoryIn, _: dict = Depends(require_admin)):
    if not await db.memory.find_one({"id": mid}):
        raise HTTPException(status_code=404, detail="Memoria no encontrada")
    updates = payload.model_dump(exclude_unset=True)
    updates["updated_at"] = now_utc()
    await db.memory.update_one({"id": mid}, {"$set": updates})
    fresh = await db.memory.find_one({"id": mid})
    return MemoryItem(**prepare(fresh, "created_at", "updated_at"))


@router.delete("/memory/{mid}")
async def delete_memory(mid: str, _: dict = Depends(require_admin)):
    await db.memory.delete_one({"id": mid})
    return {"ok": True}


# ------------------------------------------------------------ Argumentario


@router.get("/argumentario", response_model=list[Argumentario])
async def list_argumentario(model: str | None = None, _: dict = Depends(get_current_user)):
    flt = {"model": model} if model else {}
    docs = await db.argumentario.find(flt).sort("model", 1).to_list(200)
    return [Argumentario(**prepare(d, "last_updated")) for d in docs]


@router.post("/argumentario", response_model=Argumentario)
async def create_argumentario(payload: ArgumentarioIn, _: dict = Depends(require_admin)):
    obj = Argumentario(**payload.model_dump())
    await db.argumentario.insert_one(obj.model_dump())
    return obj


@router.patch("/argumentario/{aid}", response_model=Argumentario)
async def update_argumentario(aid: str, payload: ArgumentarioIn, _: dict = Depends(require_admin)):
    if not await db.argumentario.find_one({"id": aid}):
        raise HTTPException(status_code=404, detail="Argumentario no encontrado")
    updates = payload.model_dump(exclude_unset=True)
    updates["last_updated"] = now_utc()
    await db.argumentario.update_one({"id": aid}, {"$set": updates})
    fresh = await db.argumentario.find_one({"id": aid})
    return Argumentario(**prepare(fresh, "last_updated"))


@router.delete("/argumentario/{aid}")
async def delete_argumentario(aid: str, _: dict = Depends(require_admin)):
    await db.argumentario.delete_one({"id": aid})
    return {"ok": True}
